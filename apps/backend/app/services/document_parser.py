import os
from pathlib import Path
import fitz  # PyMuPDF
import docx

def extract_text_from_file(file_path: str, mime_type: str | None = None) -> str:
    path = Path(file_path)
    suffix = path.suffix.lower()
    
    if suffix == ".pdf":
        text_parts = []
        with fitz.open(file_path) as doc:
            for page in doc:
                text_parts.append(page.get_text("text"))
        return "\n".join(text_parts)
    elif suffix in [".docx", ".doc"]:
        doc = docx.Document(file_path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                paragraphs.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(paragraphs)
    elif suffix in [".txt", ".log", ".json"]:
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            return f.read()
            
    raise ValueError(f"Unsupported file format: {suffix}")
